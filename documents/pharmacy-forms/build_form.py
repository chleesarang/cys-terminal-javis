# -*- coding: utf-8 -*-
"""Yeh's Pharmacy II — Vaccine Administration Consent Form (Korean translation).
Reproduces the original layout on US Letter (8.5 x 11 in)."""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

KFONT = "맑은 고딕"
LINE = {"sz": 6, "val": "single", "color": "000000"}
NONE = {"val": "nil"}


# ---------------------------------------------------------------- primitives

TCPR_ORDER = ['w:cnfStyle','w:tcW','w:gridSpan','w:hMerge','w:vMerge','w:tcBorders','w:shd',
              'w:noWrap','w:tcMar','w:textDirection','w:tcFitText','w:vAlign','w:hideMark']
TBLPR_ORDER = ['w:tblStyle','w:tblpPr','w:tblOverlap','w:bidiVisual','w:tblStyleRowBandSize',
               'w:tblStyleColBandSize','w:tblW','w:jc','w:tblCellSpacing','w:tblInd','w:tblBorders',
               'w:shd','w:tblLayout','w:tblCellMar','w:tblLook','w:tblCaption','w:tblDescription']
TRPR_ORDER = ['w:cnfStyle','w:divId','w:gridBefore','w:gridAfter','w:wBefore','w:wAfter',
              'w:cantSplit','w:trHeight','w:tblHeader','w:tblCellSpacing','w:jc','w:hidden']


def _idx(tag, order):
    for i, n in enumerate(order):
        if qn(n) == tag:
            return i
    return None


def put(parent, el, order):
    """Remove any existing sibling with the same tag, then insert el at its schema position."""
    for child in list(parent):
        if child.tag == el.tag:
            parent.remove(child)
    my = _idx(el.tag, order)
    for child in parent:
        ci = _idx(child.tag, order)
        if ci is not None and my is not None and ci > my:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def rfonts(el, font):
    rPr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else el
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), font)


def R(p, text, size=8, bold=False, italic=False, underline=False, font=KFONT):
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.name = font
    rfonts(run._element, font)
    return run


def P(container, text="", size=8, bold=False, italic=False, align=None,
      before=0, after=0, line=1.0, indent=None, hanging=None, underline=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Inches(indent)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    if text:
        R(p, text, size=size, bold=bold, italic=italic, underline=underline)
    return p


def cell_p(cell, text="", **kw):
    """Write into a cell, reusing its mandatory first empty paragraph."""
    if cell.paragraphs and not cell.paragraphs[0].runs and not getattr(cell, "_used", False):
        p = cell.paragraphs[0]
        cell._used = True
    else:
        p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(kw.get("before", 0))
    pf.space_after = Pt(kw.get("after", 0))
    pf.line_spacing = kw.get("line", 1.0)
    if kw.get("align") is not None:
        p.alignment = kw["align"]
    if kw.get("indent") is not None:
        pf.left_indent = Inches(kw["indent"])
    if kw.get("hanging") is not None:
        pf.first_line_indent = Inches(-kw["hanging"])
    if text:
        R(p, text, size=kw.get("size", 8), bold=kw.get("bold", False),
          italic=kw.get("italic", False), underline=kw.get("underline", False))
    return p


def set_borders(cell, top=None, left=None, bottom=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for name, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if spec is None:
            continue
        el = OxmlElement("w:" + name)
        for k, v in spec.items():
            el.set(qn("w:" + k), str(v))
        tb.append(el)
    put(tcPr, tb, TCPR_ORDER)


def set_margins(cell, top=0, left=40, bottom=0, right=40):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + name)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    put(tcPr, m, TCPR_ORDER)


def vertical_text(cell, direction="btLr"):
    tcPr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), direction)
    put(tcPr, td, TCPR_ORDER)


def strip_lead(cell):
    """Drop the empty paragraph that precedes a nested table; shrink the trailing one."""
    tc = cell._tc
    kids = list(tc)
    body = [k for k in kids if k.tag in (qn('w:p'), qn('w:tbl'))]
    if body and body[0].tag == qn('w:p') and not body[0].findall(qn('w:r')):
        if any(k.tag == qn('w:tbl') for k in body):
            tc.remove(body[0])
    kids = [k for k in tc if k.tag == qn('w:p')]
    if kids:
        last = kids[-1]
        if not last.findall(qn('w:r')):
            pPr = last.find(qn('w:pPr'))
            if pPr is None:
                pPr = OxmlElement('w:pPr')
                last.insert(0, pPr)
            rPr = OxmlElement('w:rPr')
            szv = OxmlElement('w:sz'); szv.set(qn('w:val'), '2')
            rPr.append(szv)
            pPr.append(rPr)
            sp = OxmlElement('w:spacing')
            sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0')
            sp.set(qn('w:line'), '20'); sp.set(qn('w:lineRule'), 'exact')
            pPr.insert(0, sp)


def mk_table(parent, rows, cols, widths, borderless=True):
    t = parent.add_table(rows, cols) if not hasattr(parent, "_tc") else parent.add_table(rows, cols)
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    put(tblPr, lay, TBLPR_ORDER)
    if borderless:
        tb = OxmlElement("w:tblBorders")
        for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + name)
            el.set(qn("w:val"), "nil")
            tb.append(el)
        put(tblPr, tb, TBLPR_ORDER)
    for r in t.rows:
        for i, c in enumerate(r.cells):
            c.width = Inches(widths[i])
            set_margins(c, 0, 30, 0, 30)
    for i, col in enumerate(t.columns):
        col.width = Inches(widths[i])
    return t


def row_height(row, inches, exact=False):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(inches * 1440)))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    put(trPr, h, TRPR_ORDER)


# ---------------------------------------------------------------- document
doc = Document()
st = doc.styles["Normal"]
st.font.name = KFONT
st.font.size = Pt(8)
rfonts(st.element.get_or_add_rPr(), KFONT)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.0

s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.left_margin = s.right_margin = Inches(0.5)
s.top_margin = Inches(0.42)
s.bottom_margin = Inches(0.35)
W = 7.5

C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT
J = WD_ALIGN_PARAGRAPH.JUSTIFY

# ---- header ----------------------------------------------------------------
P(doc, "Yeh's Pharmacy II", size=15, bold=True, align=C)
P(doc, "966 S. Western Ave. #103, Los Angeles, CA 90006", size=9.5, align=C, before=1)
P(doc, "전화 (323) 733-7788    팩스 (323) 733-4818", size=9.5, align=C)
P(doc, "백신 접종 동의서", size=12.5, bold=True, align=C, before=4, after=4)

# ---- patient info + race/ethnicity -----------------------------------------
top = mk_table(doc, 1, 2, [4.45, 3.05])
lcell, rcell = top.rows[0].cells
set_margins(lcell, 0, 0, 0, 60)
set_margins(rcell, 0, 60, 0, 0)

FIELD_LBL = 7.5
FIELD_H = 0.17


def field_grid(cell, spec, widths, field_idx):
    """spec: list of rows, each a list of labels aligned with field_idx columns."""
    n = len(spec)
    t = mk_table(cell, n * 2, len(widths), widths)
    for ri, labels in enumerate(spec):
        line_row = t.rows[ri * 2]
        lbl_row = t.rows[ri * 2 + 1]
        row_height(line_row, FIELD_H)
        for ci in range(len(widths)):
            lc = line_row.cells[ci]
            bc = lbl_row.cells[ci]
            set_margins(lc, 0, 0, 0, 0)
            set_margins(bc, 0, 0, 20, 0)
            cell_p(lc, "", size=FIELD_LBL)
            if ci in field_idx:
                k = field_idx.index(ci)
                set_borders(lc, bottom=LINE)
                cell_p(bc, labels[k] if k < len(labels) else "", size=FIELD_LBL)
            else:
                cell_p(bc, "", size=FIELD_LBL)
    return t


field_grid(
    lcell,
    [["이름", "중간이름 이니셜", "성"],
     ["휴대전화", "생년월일 (MM/DD/YYYY)", "성별"]],
    [1.65, 0.10, 0.85, 0.10, 1.75],
    [0, 2, 4],
)

# race / ethnicity
rt = mk_table(rcell, 1, 2, [1.50, 1.55])
rc, ec = rt.rows[0].cells
set_margins(rc, 0, 0, 0, 20)
set_margins(ec, 0, 20, 0, 0)
cell_p(rc, "인종", size=8, bold=True, underline=True)
for it in ["백인", "흑인/아프리카계 미국인", "아시아계",
           "아메리카 원주민 또는 알래스카 원주민", "하와이 원주민/태평양 제도 주민"]:
    cell_p(rc, it, size=7.5, line=1.05)
cell_p(ec, "민족", size=8, bold=True, underline=True)
for it in ["히스패닉 또는 라티노", "히스패닉 또는 라티노 아님"]:
    cell_p(ec, it, size=7.5, line=1.05)

strip_lead(lcell)
strip_lead(rcell)

P(doc, "", size=3)

# full-width fields
b = mk_table(doc, 6, 3, [3.60, 0.15, 3.75])
rows_spec = [["자택 주소", "도시, 주, 우편번호"],
             ["담당 의사 이름", "담당 의사 전화번호"],
             ["알레르기", "메디케어 번호 (해당되는 경우)"]]
for ri, labels in enumerate(rows_spec):
    lr, br = b.rows[ri * 2], b.rows[ri * 2 + 1]
    row_height(lr, FIELD_H)
    for ci in range(3):
        set_margins(lr.cells[ci], 0, 0, 0, 0)
        set_margins(br.cells[ci], 0, 0, 20, 0)
        cell_p(lr.cells[ci], "", size=FIELD_LBL)
        if ci in (0, 2):
            set_borders(lr.cells[ci], bottom=LINE)
            cell_p(br.cells[ci], labels[0 if ci == 0 else 1], size=FIELD_LBL)
        else:
            cell_p(br.cells[ci], "", size=FIELD_LBL)

P(doc, "", size=4)

# ---- eligibility questions --------------------------------------------------
QUESTIONS = [
    "오늘 열이 있거나 몸이 아프십니까?",
    "지난 14일 이내에 다음 중 하나라도 겪은 적이 있습니까: 발열, 평소와 다른 기침, 평소와 다른 호흡곤란?",
    "지난 14일 이내에 본인 또는 같은 집에 사는 사람이 코로나19(COVID-19) 진단을 받았습니까?",
    "약물, 음식(예: 계란), 라텍스 또는 백신 성분(예: 소 단백질, 젤라틴, 겐타마이신, 폴리믹신, "
    "네오마이신, 페놀 또는 티메로살)에 알레르기가 있습니까? 있다면 알레르기를 적어 주십시오. ______________________",
    "지난 28일 이내에 백신 접종이나 피부 반응 검사를 받았습니까? 받았다면 백신명을 적어 주십시오. ____________",
    "과거에 독감(인플루엔자) 백신이나 다른 백신에 심각한 반응을 보인 적이 있습니까?",
    "발작약을 복용 중인 발작 질환, 뇌 질환, 길랑-바레 증후군(마비를 일으키는 질환) 또는 기타 신경계 질환이 있었던 적이 있습니까?",
    "만 65세 이상이십니까?",
    "흡연을 하십니까?",
    "__CHRONIC__",
    "8번, 9번 또는 10번 질문에 '예'라고 답하셨다면, 폐렴 백신을 접종받은 적이 있습니까?",
    "대상포진 백신을 접종받은 적이 있습니까? (만 50세 이상 환자만 해당)",
    "여성의 경우: 임신 중이거나 수유 중이거나 앞으로 한 달 이내에 임신할 계획이 있습니까?",
    "지난 3개월 동안 프레드니손 등 스테로이드제, 항암제, 류마티스 관절염·크론병·건선 치료제 등 "
    "면역체계에 영향을 주는 약을 복용했거나 방사선 치료를 받은 적이 있습니까?",
    "암, 백혈병, 림프종, HIV/에이즈 또는 기타 면역체계 질환이 있거나, 면역력이 심하게 저하된 "
    "사람과 접촉하고 있습니까?",
    "지난 1년 이내에 수혈이나 혈액제제를 투여받았거나 면역(감마) 글로불린이라는 약을 투여받은 적이 있습니까?",
    "아스피린 치료 또는 아스피린 함유 약물 치료를 받고 있습니까? (만 18세 이하만 해당)",
]

QW = [0.30, 5.70, 0.50, 0.50, 0.50]
nrows = 2 + len(QUESTIONS)
q = mk_table(doc, nrows, 5, QW, borderless=True)

for r in q.rows:
    for c in r.cells:
        set_margins(c, 10, 55, 10, 55)

# header row: Yes / No / Don't Know
hdr = q.rows[0]
row_height(hdr, 0.30)
set_borders(hdr.cells[0], top=NONE, left=NONE, bottom=NONE, right=NONE)
set_borders(hdr.cells[1], top=NONE, left=NONE, bottom=NONE, right=NONE)
cell_p(hdr.cells[0], "")
cell_p(hdr.cells[1], "")
for i, lbl in zip((2, 3, 4), ("예", "아니오", "모름")):
    c = hdr.cells[i]
    set_borders(c, top=LINE, left=LINE, bottom=LINE, right=LINE)
    c.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    cell_p(c, lbl, size=8, bold=True, align=C)

# intro row
intro = q.rows[1]
set_borders(intro.cells[0], top=NONE, left=NONE, bottom=NONE, right=NONE)
cell_p(intro.cells[0], "")
set_borders(intro.cells[1], top=LINE, left=LINE, bottom=LINE, right=LINE)
cell_p(intro.cells[1], "다음 질문들은 오늘 백신 접종이 가능한지 판단하는 데 도움이 됩니다.",
       size=8, italic=True)
for i in (2, 3, 4):
    set_borders(intro.cells[i], top=LINE, left=LINE, bottom=LINE, right=LINE)
    cell_p(intro.cells[i], "")

LIVE_START = 13  # question index (0-based) where LIVE VACCINES bracket begins -> Q14

for qi, text in enumerate(QUESTIONS):
    row = q.rows[2 + qi]
    c0, c1 = row.cells[0], row.cells[1]
    if qi < LIVE_START:
        set_borders(c0, top=NONE, left=NONE, bottom=NONE, right=NONE)
        cell_p(c0, "")
    set_borders(c1, top=LINE, left=LINE, bottom=LINE, right=LINE)
    if text == "__CHRONIC__":
        p = cell_p(c1, "", size=8, indent=0.20, hanging=0.20)
        R(p, "10.  만성 질환이나 장기적인 건강 문제가 있습니까? 있다면 해당되는 항목에 모두 동그라미 치십시오.", size=8)
        p2 = cell_p(c1, "빈혈        천식        당뇨병        심장 질환        신장 질환        "
                        "간 질환        폐 질환        기타", size=8, indent=0.42, line=1.0)
    else:
        cell_p(c1, "%d.  %s" % (qi + 1, text), size=8, indent=0.20, hanging=0.20, line=1.05)
    for i in (2, 3, 4):
        set_borders(row.cells[i], top=LINE, left=LINE, bottom=LINE, right=LINE)
        cell_p(row.cells[i], "")

# LIVE VACCINES vertical bracket over Q14..Q17
first = q.cell(2 + LIVE_START, 0)
last = q.cell(2 + len(QUESTIONS) - 1, 0)
merged = first.merge(last)
set_borders(merged, top=LINE, left=LINE, bottom=LINE, right=LINE)
set_margins(merged, 20, 10, 20, 10)
vertical_text(merged)
merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
mp = merged.paragraphs[0]
mp.alignment = C
mp.paragraph_format.space_before = Pt(0)
mp.paragraph_format.space_after = Pt(0)
R(mp, "생백신", size=7.5, bold=True)

P(doc, "", size=5)

# ---- consent statement ------------------------------------------------------
cp = P(doc, "", size=8.5, align=J, line=1.15)
R(cp, "본인은 ", size=8.5)
R(cp, "                                                      ", size=8.5, underline=True)
R(cp, " 백신을 본인 또는 본인이 대신 결정할 권한을 가진 아래에 기재된 사람에게 접종해 주실 것을 "
      "요청합니다. 본인은 백신 접종 후 최소 10분 동안 클리닉에 머무르는 데 동의합니다. 본인은 이 백신에 "
      "관한 정보를 읽었거나 설명을 들었습니다. 본인은 이 백신과 관련된 이점과 위험을 이해하며, 그 위험을 "
      "감수하기로 선택합니다. 모든 의료 처치와 마찬가지로, 백신으로 인한 이상 반응이 발생하지 않는다는 "
      "보장은 없다는 점을 인지하고 있습니다.", size=8.5)

P(doc, "", size=10)

# ---- signature --------------------------------------------------------------
sg = mk_table(doc, 2, 4, [0.85, 4.15, 0.55, 1.95])
row_height(sg.rows[0], 0.22)
r0, r1 = sg.rows[0], sg.rows[1]
r0.cells[0].vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
cell_p(r0.cells[0], "환자 서명:", size=8.5)
set_borders(r0.cells[1], bottom=LINE)
cell_p(r0.cells[1], "", size=8.5)
r0.cells[2].vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
cell_p(r0.cells[2], "날짜:", size=8.5)
set_borders(r0.cells[3], bottom=LINE)
cell_p(r0.cells[3], "", size=8.5)
cell_p(r1.cells[0], "")
cell_p(r1.cells[1], "(미성년자인 경우 부모 또는 보호자)", size=7.5, align=C)
cell_p(r1.cells[2], "")
cell_p(r1.cells[3], "")

P(doc, "", size=8)

# ---- pharmacy use only ------------------------------------------------------
P(doc, "약국 기재란", size=9, bold=True, align=C, after=2)

PW = [1.25, 0.90, 0.85, 1.15, 0.80, 0.75, 0.95, 0.85]
pt = mk_table(doc, 4, 8, PW, borderless=True)
heads = ["백신", "로트 번호", "유효기간", "제조사", "용량 (mL)", "투여 경로", "접종 부위", "VIS 날짜"]
for ri, row in enumerate(pt.rows):
    row_height(row, 0.30)
    for ci, c in enumerate(row.cells):
        set_borders(c, top=LINE, left=LINE, bottom=LINE, right=LINE)
        set_margins(c, 15, 40, 15, 40)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_p(c, heads[ci] if ri == 0 else "", size=8,
               bold=(ri == 0), align=C if ri == 0 else L)

P(doc, "", size=14)

# ---- immunizer --------------------------------------------------------------
im = mk_table(doc, 2, 5, [2.55, 0.20, 2.55, 0.20, 2.00])
row_height(im.rows[0], 0.20)
for ci in (0, 2, 4):
    set_borders(im.rows[0].cells[ci], bottom=LINE)
for ci in range(5):
    cell_p(im.rows[0].cells[ci], "", size=8)
for ci, lbl in ((0, "접종자 이름 및 직책"), (2, "접종자 서명"), (4, "접종일")):
    cell_p(im.rows[1].cells[ci], lbl, size=8)
for ci in (1, 3):
    cell_p(im.rows[1].cells[ci], "")

out = "/mnt/user-data/working/백신접종동의서_Yehs_Pharmacy_II_한글.docx"
doc.save(out)
print("saved:", out)
